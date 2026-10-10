# -*- coding: utf-8 -*-
"""
C 組測試：驗證 check_pptx.py 簡報檔結構與去 AI 味檢查工具
被測目標：scripts/check_pptx.py <pptx路徑>
注意：fixtures 在測試內動態產生，不存入二進位檔案到版控
"""

import io
import subprocess
import sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPT_PATH = str(REPO_ROOT / "scripts" / "check_pptx.py")


def _get_dummy_image_stream() -> io.BytesIO:
    """動態產生 100x100 像素之合法 PNG 二進位串流"""
    img = Image.new("RGB", (100, 100), color=(200, 200, 200))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def _create_base_presentation(total_slides: int = 8) -> Presentation:
    """動態產生合法的 PPTX 物件，每頁皆含文字框與備忘錄"""
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # 空白版型

    for i in range(total_slides):
        slide = prs.slides.add_slide(blank_layout)
        # 加入可編輯文字框
        tx_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        tf = tx_box.text_frame
        tf.text = f"第 {i + 1} 頁可編輯標題與內容"
        # 加入講者備忘錄
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = f"第 {i + 1} 頁講者口語提示與備忘"

    return prs


def test_full_slide_picture(tmp_path):
    """情境：第 5 頁只有一張滿版圖片，預期輸出包含「第 5 頁是整頁圖片，無法編輯」"""
    prs = _create_base_presentation(total_slides=8)
    # 修改第 5 頁（索引 4）：清空其形狀，改為只插入一張滿版圖片
    slide_5 = prs.slides[4]
    # 清除既有形狀
    sp_tree = slide_5.shapes._spTree
    for sp in list(slide_5.shapes):
        sp_tree.remove(sp._element)

    img_stream = _get_dummy_image_stream()
    slide_5.shapes.add_picture(img_stream, Inches(0), Inches(0), Inches(10), Inches(7.5))

    pptx_path = tmp_path / "full_picture.pptx"
    prs.save(str(pptx_path))

    result = subprocess.run(
        [sys.executable, SCRIPT_PATH, str(pptx_path)],
        capture_output=True,
        text=True,
    )
    output = result.stdout + result.stderr
    assert result.returncode != 0, "整頁滿版圖片時應該回傳非零退出碼"
    assert "第 5 頁是整頁圖片，無法編輯" in output, f"預期錯誤訊息未出現，實際輸出：\n{output}"


def test_missing_notes(tmp_path):
    """情境：第 7 頁備忘錄空白，預期輸出包含「第 7 頁沒有講者備忘錄」"""
    prs = _create_base_presentation(total_slides=8)
    # 修改第 7 頁（索引 6）：清空備忘錄
    slide_7 = prs.slides[6]
    slide_7.notes_slide.notes_text_frame.text = ""

    pptx_path = tmp_path / "missing_notes.pptx"
    prs.save(str(pptx_path))

    result = subprocess.run(
        [sys.executable, SCRIPT_PATH, str(pptx_path)],
        capture_output=True,
        text=True,
    )
    output = result.stdout + result.stderr
    assert result.returncode != 0, "缺少講者備忘錄時應該回傳非零退出碼"
    assert "第 7 頁沒有講者備忘錄" in output, f"預期錯誤訊息未出現，實際輸出：\n{output}"


def test_no_text_frame(tmp_path):
    """情境：第 2 頁沒有文字框，預期輸出包含「第 2 頁沒有可編輯的文字框」"""
    prs = _create_base_presentation(total_slides=8)
    # 修改第 2 頁（索引 1）：清空文字框，只放無文字的純矩形圖案
    slide_2 = prs.slides[1]
    sp_tree = slide_2.shapes._spTree
    for sp in list(slide_2.shapes):
        sp_tree.remove(sp._element)

    # 插入一個空矩形形狀，無文字
    shape = slide_2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    # 確認 shape 的 text_frame 為空或非可編輯文字框
    shape.text_frame.text = ""

    pptx_path = tmp_path / "no_text_frame.pptx"
    prs.save(str(pptx_path))

    result = subprocess.run(
        [sys.executable, SCRIPT_PATH, str(pptx_path)],
        capture_output=True,
        text=True,
    )
    output = result.stdout + result.stderr
    assert result.returncode != 0, "缺少可編輯文字框時應該回傳非零退出碼"
    assert "第 2 頁沒有可編輯的文字框" in output, f"預期錯誤訊息未出現，實際輸出：\n{output}"


def test_valid_pptx_passes(tmp_path):
    """情境：每頁有文字框與備忘錄，預期 exit 0"""
    prs = _create_base_presentation(total_slides=8)
    pptx_path = tmp_path / "valid.pptx"
    prs.save(str(pptx_path))

    result = subprocess.run(
        [sys.executable, SCRIPT_PATH, str(pptx_path)],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, f"合法 PPTX 檢查應通過，但退出碼為 {result.returncode}，輸出：\n{result.stdout}\n{result.stderr}"
