# -*- coding: utf-8 -*-
"""
測試：驗證 check_grounding.py 投影片內容真實性與數字同源檢查工具
被測目標：scripts/check_grounding.py <pptx路徑> <ssot路徑>
"""

import subprocess
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPT_PATH = str(REPO_ROOT / "scripts" / "check_grounding.py")

SAMPLE_SSOT_CONTENT = """# 小狗日常 BNI 專題簡報 — 中繼 Markdown (SSOT)

## 0. 核心定位
- 核心主張：狗狗洗完澡不會發抖，是因為我們先陪牠玩 10 分鐘
- 理想引薦對象：住在東區 5 公里內、家裡狗狗一看到外出籠就躲到床底下的飼主

## 1. 投影片結構規劃 (Slide Outline)
- P01 封面：狗狗洗完澡不會發抖，是因為我們先陪牠玩 10 分鐘
- P02 現狀卡關：飼主要趴在地上用零食把狗狗哄出來，每次花 20 分鐘
- P05 核心機制：先陪玩 10 分鐘，再用低噪音吹風機分段吹乾
- P06 交付標準：洗完傳 30 秒影片，讓飼主看到狗狗沒發抖；到府接送：東區 5 公里內免費
"""


def _create_presentation_with_text(slides_text: list[str]) -> Presentation:
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    for text in slides_text:
        slide = prs.slides.add_slide(blank_layout)
        tx_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        tx_box.text_frame.text = text
        slide.notes_slide.notes_text_frame.text = "備忘提示"
    return prs


def test_invented_percentage_rejected(tmp_path):
    """情境：投影片包含捏造百分比「噪音降低 45%」，SSOT 中無 45，預期報錯失敗"""
    ssot_path = tmp_path / "ssot.md"
    ssot_path.write_text(SAMPLE_SSOT_CONTENT, encoding="utf-8")

    prs = _create_presentation_with_text([
        "狗狗洗完澡不會發抖，是因為我們先陪牠玩 10 分鐘",
        "採用低噪音吹風機，噪音降低 45%"
    ])
    pptx_path = tmp_path / "invented_percentage.pptx"
    prs.save(str(pptx_path))

    result = subprocess.run(
        [sys.executable, SCRIPT_PATH, str(pptx_path), str(ssot_path)],
        capture_output=True,
        text=True,
    )
    output = result.stdout
    assert result.returncode != 0, f"捏造百分比應該回傳非零退出碼，實際輸出：\n{output}"
    assert "45" in output, f"預期錯誤應指出捏造數字 45，實際輸出：\n{output}"


def test_invented_minutes_rejected(tmp_path):
    """情境：投影片包含捏造分鐘數「到府前 15 分鐘 LINE 通知」或「每 3 分鐘稍作撫觸」，SSOT 無 15 或 3，預期失敗"""
    ssot_path = tmp_path / "ssot.md"
    ssot_path.write_text(SAMPLE_SSOT_CONTENT, encoding="utf-8")

    prs = _create_presentation_with_text([
        "到府前 15 分鐘 LINE 通知",
        "吹風過程每 3 分鐘稍作撫觸"
    ])
    pptx_path = tmp_path / "invented_minutes.pptx"
    prs.save(str(pptx_path))

    result = subprocess.run(
        [sys.executable, SCRIPT_PATH, str(pptx_path), str(ssot_path)],
        capture_output=True,
        text=True,
    )
    output = result.stdout
    assert result.returncode != 0, f"捏造分鐘數應該回傳非零退出碼，實際輸出：\n{output}"
    assert ("15" in output or "3" in output), f"預期錯誤應指出捏造數字 15 或 3，實際輸出：\n{output}"


def test_grounded_numbers_pass(tmp_path):
    """情境：投影片數字皆存在於 SSOT（10 分鐘、20 分鐘、30 秒、5 公里），預期通過 exit 0"""
    ssot_path = tmp_path / "ssot.md"
    ssot_path.write_text(SAMPLE_SSOT_CONTENT, encoding="utf-8")

    prs = _create_presentation_with_text([
        "先陪牠玩 10 分鐘，每次花 20 分鐘哄出來",
        "洗完傳 30 秒影片，東區 5 公里內免費"
    ])
    pptx_path = tmp_path / "grounded.pptx"
    prs.save(str(pptx_path))

    result = subprocess.run(
        [sys.executable, SCRIPT_PATH, str(pptx_path), str(ssot_path)],
        capture_output=True,
        text=True,
    )
    output = result.stdout + result.stderr
    assert result.returncode == 0, f"SSOT 內真實數字應檢查通過，但退出碼為 {result.returncode}，輸出：\n{output}"


def test_page_numbers_ignored(tmp_path):
    """情境：投影片包含頁碼標記如 P05 / 08 或 01，不被誤判為捏造數字，預期通過 exit 0"""
    ssot_path = tmp_path / "ssot.md"
    ssot_path.write_text(SAMPLE_SSOT_CONTENT, encoding="utf-8")

    prs = _create_presentation_with_text([
        "P01 / 08 封面破題大金句",
        "P05 / 08 先陪玩 10 分鐘"
    ])
    pptx_path = tmp_path / "page_numbers.pptx"
    prs.save(str(pptx_path))

    result = subprocess.run(
        [sys.executable, SCRIPT_PATH, str(pptx_path), str(ssot_path)],
        capture_output=True,
        text=True,
    )
    output = result.stdout + result.stderr
    assert result.returncode == 0, f"頁碼標記不應觸發捏造數字報錯，退出碼為 {result.returncode}，輸出：\n{output}"
